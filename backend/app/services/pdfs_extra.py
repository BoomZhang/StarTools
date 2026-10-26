"""PDF P1：页面管理、压缩、PDF 转 Word、密码；Office 转 PDF。"""

import subprocess
from pathlib import Path

import pymupdf as fitz
from docx import Document as DocxDocument
from pypdf import PdfReader, PdfWriter

from app.services.common import (
    expand_ranges,
    normalize_ranges,
    safe_stem,
    unique_filename,
)
from app.services.types import FileResult, ProcessOutcome, TaskError

OFFICE_SUFFIXES = {".docx", ".xlsx", ".pptx"}
SOFFICE_CANDIDATES = (
    "/usr/bin/soffice",
    "/usr/local/bin/soffice",
    "/Applications/LibreOffice.app/Contents/MacOS/soffice",
)


def office_available() -> bool:
    import shutil

    if shutil.which("soffice") or shutil.which("libreoffice"):
        return True
    return any(Path(path).exists() for path in SOFFICE_CANDIDATES)


def _soffice_binary() -> str | None:
    import shutil

    for name in ("soffice", "libreoffice"):
        path = shutil.which(name)
        if path:
            return path
    for candidate in SOFFICE_CANDIDATES:
        if Path(candidate).exists():
            return candidate
    return None


def _reader(path: Path, name: str) -> PdfReader:
    try:
        reader = PdfReader(str(path))
    except Exception as exc:
        raise TaskError("E1-400", f"{name} 无法读取，文件可能已损坏") from exc
    if reader.is_encrypted:
        raise TaskError("E1-400", f"{name} 已加密，请先移除打开密码")
    if len(reader.pages) == 0:
        raise TaskError("E1-400", f"{name} 不含任何页面")
    return reader


def _parse_pages(expression: str, total: int, label: str) -> list[int]:
    if not expression.strip():
        return []
    try:
        parts = normalize_ranges(expression)
    except ValueError as exc:
        raise TaskError("E1-400", f"{label}无法解析：{exc}；示例 1-3,5,8-10") from exc
    try:
        pages = expand_ranges(parts, total)
    except ValueError:
        raise TaskError("E2-413", f"{label}超出有效范围 1–{total}", 413) from None
    return pages


def process_pdf_pages(
    files: list[Path], names: list[str], params: dict, workdir: Path
) -> ProcessOutcome:
    path = files[0]
    name = names[0]
    reader = _reader(path, name)
    total = len(reader.pages)
    if total > 500:
        raise TaskError("E2-413", f"PDF 共 {total} 页，超过 500 页上限", 413)

    rotate = int(params.get("rotate", 0) or 0)
    if rotate not in {0, 90, 180, 270}:
        raise TaskError("E1-400", "旋转角度仅支持 0 / 90 / 180 / 270")
    remove = set(_parse_pages(str(params.get("remove", "")), total, "删除页码范围"))
    extract = _parse_pages(str(params.get("extract", "")), total, "抽取页码范围")
    order = _parse_pages(str(params.get("order", "")), total, "页面顺序")

    if extract:
        pages = extract
    elif order:
        pages = order
    else:
        pages = list(range(1, total + 1))
    pages = [page for page in pages if page not in remove]
    if not pages:
        raise TaskError("E1-400", "删除与抽取后没有剩余页面")

    writer = PdfWriter()
    for page_number in pages:
        page = reader.pages[page_number - 1]
        if rotate:
            page.rotate(rotate)
        writer.add_page(page)

    target = workdir / unique_filename(f"{safe_stem(name)}_页面管理.pdf", set())
    with target.open("wb") as handle:
        writer.write(handle)
    return ProcessOutcome(
        results=[
            FileResult(
                filename=target.name,
                path=target,
                content_type="application/pdf",
                meta={"pages": len(pages), "rotate": rotate},
            )
        ],
        notes=["页面顺序按抽取或自定义顺序输出；原书签、表单与签名不保留。"],
    )


def process_pdf_compress(
    files: list[Path], names: list[str], params: dict, workdir: Path
) -> ProcessOutcome:
    path = files[0]
    name = names[0]
    level = str(params.get("level", "balanced"))
    settings = {
        "light": {"dpi": 220, "quality": 80, "effort": 60},
        "balanced": {"dpi": 150, "quality": 60, "effort": 80},
        "max": {"dpi": 100, "quality": 40, "effort": 100},
    }.get(level, {"dpi": 150, "quality": 60, "effort": 80})

    try:
        document = fitz.open(str(path))
    except Exception as exc:
        raise TaskError("E1-400", "PDF 无法读取，文件可能已损坏") from exc
    if document.needs_pass or document.is_encrypted:
        document.close()
        raise TaskError("E1-400", "PDF 已加密，请先移除打开密码")
    if document.page_count > 500:
        document.close()
        raise TaskError("E2-413", "PDF 超过 500 页上限", 413)

    original_size = path.stat().st_size
    target = workdir / unique_filename(f"{safe_stem(name)}_压缩.pdf", set())
    try:
        document.rewrite_images(
            dpi_threshold=settings["dpi"] + 50,
            dpi_target=settings["dpi"],
            quality=settings["quality"],
            lossy=True,
            lossless=True,
        )
        document.save(
            str(target),
            garbage=4,
            clean=True,
            deflate=True,
            deflate_images=True,
            deflate_fonts=True,
            compression_effort=settings["effort"],
        )
    finally:
        document.close()

    new_size = target.stat().st_size
    note = None
    if new_size >= original_size:
        target.unlink(missing_ok=True)
        target = workdir / unique_filename(name, set())
        target.write_bytes(path.read_bytes())
        new_size = original_size
        note = "压缩后未更小，已保留原文件。"

    return ProcessOutcome(
        results=[
            FileResult(
                filename=target.name,
                path=target,
                content_type="application/pdf",
                meta={
                    "original_size": original_size,
                    "new_size": new_size,
                    "ratio": round((1 - new_size / original_size) * 100, 1) if original_size else 0,
                    "level": level,
                    "note": note,
                },
            )
        ],
        notes=["仅对嵌入图片降采样，文字仍保持可选中；原签名失效。"],
    )


def process_pdf_to_word(
    files: list[Path], names: list[str], params: dict, workdir: Path
) -> ProcessOutcome:
    path = files[0]
    name = names[0]
    try:
        document = fitz.open(str(path))
    except Exception as exc:
        raise TaskError("E1-400", "PDF 无法读取，文件可能已损坏") from exc
    if document.needs_pass or document.is_encrypted:
        document.close()
        raise TaskError("E1-400", "PDF 已加密，请先移除打开密码")
    if document.page_count > 500:
        document.close()
        raise TaskError("E2-413", "PDF 超过 500 页上限", 413)

    docx = DocxDocument()
    total_text = 0
    for index, page in enumerate(document):
        blocks = page.get_text("blocks")
        blocks.sort(key=lambda block: (round(block[1], 1), block[0]))
        for block in blocks:
            text = str(block[4] or "").strip()
            if not text:
                continue
            total_text += len(text)
            paragraph = docx.add_paragraph()
            run = paragraph.add_run(text)
            run.font.name = "Microsoft YaHei"
        if index < document.page_count - 1:
            docx.add_page_break()
    document.close()

    if total_text < 20:
        raise TaskError("E3-422", "没有提取到可用文本，扫描件请改用 OCR-02")

    target = workdir / unique_filename(f"{safe_stem(name)}.docx", set())
    docx.save(str(target))
    return ProcessOutcome(
        results=[
            FileResult(
                filename=target.name,
                path=target,
                content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                meta={"chars": total_text},
            )
        ],
        notes=["仅支持文本型 PDF；版式为尽力保留，不承诺与原文件 100% 一致。"],
    )


def process_pdf_password(
    files: list[Path], names: list[str], params: dict, workdir: Path
) -> ProcessOutcome:
    path = files[0]
    name = names[0]
    mode = str(params.get("mode", "add"))
    password = str(params.get("password", ""))
    if not password or len(password) > 64:
        raise TaskError("E1-400", "密码需在 1–64 个字符之间")

    reader = _reader(path, name) if mode == "add" else PdfReader(str(path))
    writer = PdfWriter()

    if mode == "add":
        writer.append(reader)
        writer.encrypt(password, algorithm="AES-256")
        output_name = f"{safe_stem(name)}_加密.pdf"
        note = "请妥善保存密码；遗失后无法恢复。"
    elif mode == "remove":
        if not reader.is_encrypted:
            raise TaskError("E1-400", "该 PDF 未加密，无需移除密码")
        result = reader.decrypt(password)
        if not result:
            raise TaskError("E1-400", "密码不正确，无法移除保护")
        writer.append(reader)
        output_name = f"{safe_stem(name)}_已解密.pdf"
        note = "已移除打开密码；原签名失效。"
    else:
        raise TaskError("E1-400", "未知的密码操作")

    target = workdir / unique_filename(output_name, set())
    with target.open("wb") as handle:
        writer.write(handle)
    return ProcessOutcome(
        results=[
            FileResult(
                filename=target.name,
                path=target,
                content_type="application/pdf",
                meta={"mode": mode},
            )
        ],
        notes=[note, "密码不写入日志、使用记录或操作记录。"],
    )


def process_office_to_pdf(
    files: list[Path], names: list[str], params: dict, workdir: Path
) -> ProcessOutcome:
    path = files[0]
    name = names[0]
    suffix = path.suffix.lower()
    if suffix not in OFFICE_SUFFIXES:
        raise TaskError("E1-400", "仅支持 DOCX、XLSX、PPTX")
    binary = _soffice_binary()
    if not binary:
        raise TaskError("E4-503", "Office 转换引擎不可用，请联系维护者", 503)

    profile = workdir / ".lo-profile"
    profile.mkdir(exist_ok=True)
    command = [
        binary,
        "--headless",
        "--norestore",
        f"-env:UserInstallation=file://{profile}",
        "--convert-to",
        "pdf",
        "--outdir",
        str(workdir),
        str(path),
    ]
    try:
        result = subprocess.run(command, capture_output=True, text=True, timeout=180)
    except subprocess.TimeoutExpired as exc:
        raise TaskError("E4-504", "Office 转换超时，请稍后重试", 504) from exc
    if result.returncode != 0:
        detail = (result.stderr or result.stdout or "").strip()[:200]
        raise TaskError("E3-422", f"Office 转换失败：{detail or '未知原因'}")

    candidates = [
        item for item in workdir.glob("*.pdf") if item.name != path.name and item.name.endswith(".pdf")
    ]
    if not candidates:
        raise TaskError("E3-422", "转换未生成 PDF 文件")
    produced = candidates[0]
    target = workdir / unique_filename(f"{safe_stem(name)}.pdf", set())
    produced.replace(target)
    return ProcessOutcome(
        results=[
            FileResult(
                filename=target.name,
                path=target,
                content_type="application/pdf",
                meta={"source": suffix.lstrip(".")},
            )
        ],
        notes=["字体缺失可能导致版式变化；宏与嵌入脚本不会执行。"],
    )
