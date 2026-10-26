"""OCR P1：PDF 逐页识别与表格识别导出 XLSX。"""

from pathlib import Path

import pymupdf as fitz
import pytesseract
from openpyxl import Workbook
from openpyxl.styles import Font
from PIL import Image, ImageOps
from pytesseract import Output

from app.services.common import (
    expand_ranges,
    normalize_ranges,
    safe_stem,
    tesseract_available,
    timestamp_label,
    unique_filename,
)
from app.services.ocr import LANGUAGE_MAP, _auto_rotate
from app.services.types import FileResult, ProcessOutcome, TaskError, TextResult

MAX_PDF_PAGES = 50


def process_pdf_ocr(
    files: list[Path], names: list[str], params: dict, workdir: Path
) -> ProcessOutcome:
    if not tesseract_available():
        raise TaskError("E4-503", "服务端 OCR 引擎暂不可用，请联系维护者", 503)
    path = files[0]
    name = names[0]
    try:
        document = fitz.open(str(path))
    except Exception as exc:
        raise TaskError("E1-400", "PDF 无法读取，文件可能已损坏") from exc
    if document.needs_pass or document.is_encrypted:
        document.close()
        raise TaskError("E1-400", "PDF 已加密，请先移除打开密码")
    total = document.page_count
    if total == 0:
        document.close()
        raise TaskError("E1-400", "PDF 不含任何页面")
    if total > MAX_PDF_PAGES:
        document.close()
        raise TaskError("E2-413", f"PDF 共 {total} 页，单次最多识别 {MAX_PDF_PAGES} 页", 413)

    expression = str(params.get("ranges", "")).strip()
    if not expression or expression == "all":
        pages = list(range(1, total + 1))
    else:
        try:
            parts = normalize_ranges(expression)
            pages = expand_ranges(parts, total)
        except ValueError:
            document.close()
            raise TaskError("E1-400", "页码范围无法解析或超出有效范围，示例 1-3,5", 400) from None

    language = LANGUAGE_MAP.get(str(params.get("language", "mixed")), "chi_sim+eng")
    auto_rotate = bool(params.get("auto_rotate", True))
    outcome = ProcessOutcome()
    stem = safe_stem(name)
    zoom = 200 / 72
    matrix = fitz.Matrix(zoom, zoom)

    for page_number in pages:
        try:
            page = document.load_page(page_number - 1)
            pixmap = page.get_pixmap(matrix=matrix, alpha=False)
            image = Image.frombytes("RGB", (pixmap.width, pixmap.height), pixmap.samples)
            if auto_rotate:
                image = _auto_rotate(image)
            text = pytesseract.image_to_string(image, lang=language, timeout=60).strip()
            if not text:
                outcome.failures.append(
                    {
                        "filename": f"第 {page_number} 页",
                        "code": "E3-422",
                        "message": f"第 {page_number} 页未识别到文字。",
                    }
                )
                continue
            outcome.results.append(
                TextResult(
                    filename=f"{stem}_p{page_number:03d}.txt",
                    text=text,
                    meta={"page": page_number, "language": language, "chars": len(text)},
                )
            )
        except Exception as exc:  # pragma: no cover
            outcome.failures.append(
                {"filename": f"第 {page_number} 页", "code": "E3-422", "message": f"第 {page_number} 页识别失败：{exc}"}
            )

    document.close()
    if not outcome.results:
        raise TaskError("E3-422", "所有页面均未识别到文字")
    outcome.notes.append("按页输出文本；识别结果需人工校对。")
    return outcome


def _tsv_words(image: Image.Image, language: str) -> list[dict]:
    data = pytesseract.image_to_data(
        image, lang=language, config="--psm 11", output_type=Output.DICT, timeout=60
    )
    words: list[dict] = []
    for index, text in enumerate(data.get("text", [])):
        value = str(text or "").strip()
        if not value:
            continue
        try:
            confidence = float(data["conf"][index])
        except (TypeError, ValueError):
            confidence = 0
        if confidence < 35:
            continue
        left = int(data["left"][index])
        top = int(data["top"][index])
        width = int(data["width"][index])
        height = int(data["height"][index])
        words.append(
            {
                "text": value,
                "left": left,
                "top": top,
                "width": width,
                "height": height,
                "cx": left + width / 2,
                "cy": top + height / 2,
            }
        )
    return words


def _rows_and_columns(words: list[dict], image_width: int) -> list[list[str]]:
    if not words:
        return []
    heights = sorted(word["height"] for word in words)
    median_height = heights[len(heights) // 2] or 12
    words = sorted(words, key=lambda word: (word["cy"], word["cx"]))

    rows: list[list[dict]] = []
    for word in words:
        placed = False
        for row in rows:
            row_center = sum(item["cy"] for item in row) / len(row)
            if abs(word["cy"] - row_center) <= median_height * 0.65:
                row.append(word)
                placed = True
                break
        if not placed:
            rows.append([word])
    for row in rows:
        row.sort(key=lambda word: word["left"])

    centers = sorted(word["cx"] for word in words)
    tolerance = max(24, image_width * 0.025)
    columns: list[float] = []
    for center in centers:
        if not columns or center - columns[-1] > tolerance:
            columns.append(center)
        else:
            columns[-1] = (columns[-1] + center) / 2

    table: list[list[str]] = []
    for row in rows:
        cells = [""] * len(columns)
        for word in row:
            best = min(range(len(columns)), key=lambda index: abs(columns[index] - word["cx"]))
            cells[best] = (cells[best] + " " + word["text"]).strip()
        table.append(cells)
    return table


def process_table_ocr(
    files: list[Path], names: list[str], params: dict, workdir: Path
) -> ProcessOutcome:
    if not tesseract_available():
        raise TaskError("E4-503", "服务端 OCR 引擎暂不可用，请联系维护者", 503)
    language = LANGUAGE_MAP.get(str(params.get("language", "mixed")), "chi_sim+eng")
    auto_rotate = bool(params.get("auto_rotate", True))
    outcome = ProcessOutcome()

    workbook = Workbook()
    workbook.remove(workbook.active)
    preview: list[list[str]] = []
    total_rows = 0

    for index, (path, name) in enumerate(zip(files, names)):
        try:
            image = Image.open(path)
            image.seek(0)
            image.load()
            image = ImageOps.exif_transpose(image).convert("RGB")
            if auto_rotate:
                image = _auto_rotate(image)
            words = _tsv_words(image, language)
            table = _rows_and_columns(words, image.width)
            if not table:
                outcome.failures.append(
                    {
                        "filename": name,
                        "code": "E3-422",
                        "message": "未识别到表格文字，请确认图片清晰度与语言。",
                    }
                )
                continue
            sheet_name = (safe_stem(name) or f"Sheet{index + 1}")[:31]
            sheet = workbook.create_sheet(title=sheet_name)
            for row_index, row in enumerate(table, start=1):
                for column_index, value in enumerate(row, start=1):
                    cell = sheet.cell(row=row_index, column=column_index, value=value)
                    if row_index == 1:
                        cell.font = Font(bold=True)
            for column_index in range(1, len(table[0]) + 1):
                sheet.column_dimensions[chr(64 + min(column_index, 26))].width = 18
            sheet.freeze_panes = "A2"
            total_rows += len(table)
            if index == 0:
                preview = table[:8]
        except RuntimeError as exc:
            outcome.failures.append({"filename": name, "code": "E4-504", "message": f"识别超时：{exc}"})
        except Exception as exc:  # pragma: no cover
            outcome.failures.append({"filename": name, "code": "E3-422", "message": f"{name} 识别失败：{exc}"})

    if not workbook.sheetnames:
        raise TaskError("E3-422", "所有图片均未识别到表格内容")

    filename = unique_filename(f"表格识别_{timestamp_label()}.xlsx", set())
    target = workdir / filename
    workbook.save(target)
    return ProcessOutcome(
        results=[
            FileResult(
                filename=filename,
                path=target,
                content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                meta={
                    "rows": total_rows,
                    "columns": len(preview[0]) if preview else 0,
                    "sheets": len(workbook.sheetnames),
                    "preview": preview,
                },
            )
        ],
        notes=["表格结构为尽力还原，合并单元格可能需要人工调整；识别结果需校对。"],
    )
